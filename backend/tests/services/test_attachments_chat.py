import pytest
import base64
import io
import pypdf
import docx
import openpyxl
from app.services.chat_service import chat_service

def test_parse_text_attachment():
    raw_txt = "Testo di prova commessa PRJ-999"
    b64 = base64.b64encode(raw_txt.encode("utf-8")).decode("utf-8")
    att = {
        "name": "note.txt",
        "type": "text/plain",
        "data": f"data:text/plain;base64,{b64}"
    }
    images, docs_text = chat_service._parse_attachments([att])
    assert len(images) == 0
    assert "PRJ-999" in docs_text
    assert "note.txt" in docs_text

def test_parse_image_attachment():
    fake_img = b"\x89PNG\r\n\x1a\nfake"
    b64 = base64.b64encode(fake_img).decode("utf-8")
    att = {
        "name": "foto_schema.png",
        "type": "image/png",
        "data": f"data:image/png;base64,{b64}"
    }
    images, docs_text = chat_service._parse_attachments([att])
    assert len(images) == 1
    assert images[0]["name"] == "foto_schema.png"
    assert images[0]["data_url"].startswith("data:image/png;base64,")

def test_parse_docx_attachment():
    doc = docx.Document()
    doc.add_paragraph("Specifiche tecniche motore brushless")
    buf = io.BytesIO()
    doc.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
    att = {
        "name": "specifiche.docx",
        "type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "data": b64
    }
    images, docs_text = chat_service._parse_attachments([att])
    assert len(images) == 0
    assert "motore brushless" in docs_text

def test_parse_excel_attachment():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Fase"
    ws["B1"] = "Addetto"
    ws.append(["Montaggio", "Mario"])
    buf = io.BytesIO()
    wb.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
    att = {
        "name": "pianificazione.xlsx",
        "type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "data": b64
    }
    images, docs_text = chat_service._parse_attachments([att])
    assert len(images) == 0
    assert "Montaggio" in docs_text
    assert "Mario" in docs_text
